# -*- coding: utf-8 -*-
"""Dựng file .xlsx báo cáo từ template tháng trước + recalc bằng Excel COM.

Vì sao phải nhờ Excel: 2 sheet Rút gọn dùng ARRAY FORMULA (INDEX/SMALL/IF) để kéo
dữ liệu từ Data_Tổng. openpyxl chỉ ghi được công thức chứ không tính được, mà
LibreOffice thì máy không có -> dùng COM. Không có Excel thì file vẫn ghi ra bình
thường, chỉ là giá trị sẽ hiện khi bạn mở file lên (Excel tự tính lúc mở).
"""
from __future__ import annotations

import os
import shutil
import time

import openpyxl
from openpyxl.utils import get_column_letter as gcl
from openpyxl.worksheet.formula import ArrayFormula

from . import rules as R

SUFFIX = '_Rút gọn'
DATA_SHEET = 'Data_Tổng'

TAG_ORDER = ['QUAYCHUPMKT', 'DESIGN MKT', 'VIDEO MKT', 'RETOUCH', 'DESIGN LX', 'VIDEO LX']
CATEGORY = {'QUAYCHUPMKT': 'Quay/Chụp', 'DESIGN MKT': 'Design', 'VIDEO MKT': 'Video social',
            'RETOUCH': 'Retouch', 'DESIGN LX': 'Design', 'VIDEO LX': 'Video SX'}

# cột B..R của Data_Tổng
RAW_ORDER = ['ngaytao', 'nguoitao', 'giaocho', 'ten', 'mota', 'theodoi', 'giaidoan',
             'trangthai', 'nhan', 'lydo', 'tghoanthanh', 'tgcapnhat', 'ngaymuon',
             'brief', 'phongban', 'kichthuoc', 'mau']

# cột sheet Rút gọn -> cột nguồn ở Data_Tổng
COLMAP = {2: 'E', 3: 'V', 4: 'W', 5: 'X', 6: 'Y', 7: 'Z', 8: 'AC', 9: 'T', 11: 'AB'}


def _find_config_col(ws) -> int:
    """Dò cột chứa nhãn 'Công chuẩn' (mặc định O = 15)."""
    for c in range(14, 21):
        for r in range(2, 7):
            if str(ws.cell(r, c).value or '').strip().lower().startswith('công chuẩn'):
                return c
    return 15


def build(tasks, template_path, out_path, persons_cfg, tag_order=None):
    """tasks: list dict đã lọc (chỉ những cái include=True).
    persons_cfg: {person: {'cong_chuan','timeoff','ot'}}
    """
    order = tag_order or TAG_ORDER
    rows = sorted([t for t in tasks if t.get('include')],
                  key=lambda t: order.index(t['tag']) if t['tag'] in order else len(order))

    shutil.copyfile(template_path, out_path)
    wb = openpyxl.load_workbook(out_path)
    warnings = []

    if DATA_SHEET not in wb.sheetnames:
        raise ValueError(f'Template thiếu sheet {DATA_SHEET}')

    # ghi nhớ vùng dữ liệu CŨ của template để còn dọn định dạng thừa bên dưới
    old_rows = {s: wb[s].max_row for s in wb.sheetnames}

    d = wb[DATA_SHEET]
    for r in range(2, d.max_row + 1):
        for c in range(1, 33):
            d.cell(r, c).value = None

    for k, t in enumerate(rows):
        r = k + 2
        raw = t.get('_raw') or {}
        d.cell(r, 1).value = t['tag']
        for j, f in enumerate(RAW_ORDER):
            v = raw.get(f)
            d.cell(r, 2 + j).value = v if v not in ('', None) else None
        d.cell(r, 19).value = t['person']
        d.cell(r, 20).value = float(t['T'])
        d.cell(r, 21).value = raw.get('tdone') or None
        for j, v in enumerate(t['ratings']):
            d.cell(r, 22 + j).value = v
        d.cell(r, 27).value = raw.get('feedback') or None
        d.cell(r, 28).value = t['id']
        d.cell(r, 29).value = f'=SUM(V{r}:Z{r})'
        d.cell(r, 30).value = CATEGORY.get(t['tag'], t['tag'])

    last_data = len(rows) + 1
    block = [(p, f'=SUMIF($S$2:$S${last_data},"{p}",$T$2:$T${last_data})')
             for p in persons_cfg]
    block.append(('Tổng thời gian phân bổ', f'=SUM($T$2:$T${last_data})'))
    for i in range(3, 3 + max(3, len(block))):
        d.cell(i, 31).value = None
        d.cell(i, 32).value = None
    for i, (lab, formula) in enumerate(block, start=3):
        d.cell(i, 31).value = lab
        d.cell(i, 32).value = formula

    summary = {}
    for person, cfg in persons_cfg.items():
        sheet = f'{person}{SUFFIX}'
        if sheet not in wb.sheetnames:
            warnings.append(f'Template không có sheet "{sheet}" — bỏ qua {person}.')
            continue
        ws = wb[sheet]
        for r in range(3, ws.max_row + 1):
            for c in range(1, 15):
                ws.cell(r, c).value = None

        sub = [t for t in rows if t['person'] == person]
        for k, t in enumerate(sub):
            r = k + 3
            ws.cell(r, 1).value = k + 1
            for col, dcol in COLMAP.items():
                f = (f'=IFERROR(INDEX({DATA_SHEET}!${dcol}:${dcol}, '
                     f'SMALL(IF({DATA_SHEET}!$S$2:$S$997="{person}", '
                     f'ROW({DATA_SHEET}!$S$2:$S$997)), ROW(A{k + 1}))), "")')
                ws.cell(r, col).value = ArrayFormula(f'{gcl(col)}{r}', f)
            ws.cell(r, 10).value = f'="https://workflow.base.vn/job/" & K{r}'
            ws.cell(r, 12).value = t['L']
            ws.cell(r, 13).value = t['M']
            ws.cell(r, 14).value = float(t['N'])

        last = len(sub) + 2
        cc = _find_config_col(ws)
        vc, qc = cc + 1, cc + 2
        ws.cell(3, cc).value = 'Công chuẩn'; ws.cell(3, vc).value = float(cfg['cong_chuan'])
        ws.cell(4, cc).value = 'TIMEOFF';    ws.cell(4, vc).value = float(cfg.get('timeoff', 0))
        ws.cell(5, cc).value = 'OT';         ws.cell(5, vc).value = float(cfg.get('ot', 0))
        ws.cell(8, cc).value = 'Điểm thời gian'
        ws.cell(8, vc).value = 'Điểm chất lượng'
        ws.cell(8, qc).value = 'Điểm KPI'
        L_cc, L_vc = gcl(cc), gcl(vc)
        ws.cell(9, cc).value = f'=((SUM(N3:N{last}))/(({L_vc}3-{L_vc}4)*8))*10'
        ws.cell(9, vc).value = f'=AVERAGE(H3:H{last})'
        ws.cell(9, qc).value = f'=AVERAGE({L_cc}9:{L_vc}9)'
        summary[person] = {'count': len(sub), 'last_row': last, 'config_col': cc,
                           'total_n': round(sum(float(t['N']) for t in sub), 2)}

    wb.save(out_path)
    sum_t = {}
    for person in persons_cfg:
        sum_t[person] = round(sum(float(t['T']) for t in rows if t['person'] == person), 2)
    return {'out_path': out_path, 'data_last_row': last_data,
            'task_count': len(rows), 'per_person': summary, 'warnings': warnings,
            '_sum_t': sum_t, '_old_rows': old_rows}


# --------------------------------------------------------------- Excel COM

def recalc(out_path, layout, template_last_rows=None):
    """Mở bằng Excel, chép định dạng xuống dòng mới, tính lại, lưu, đọc kết quả.
    Trả về dict kết quả + kiểm tra. Không có Excel -> trả available=False."""
    try:
        import win32com.client as w32
    except ImportError:
        return {'available': False, 'reason': 'Máy không có pywin32 / Excel.'}

    xlPasteFormats = -4122
    app = None
    try:
        app = w32.DispatchEx('Excel.Application')
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Open(os.path.abspath(out_path))

        pivots = []
        for ws in wb.Worksheets:
            try:
                n = ws.PivotTables().Count
            except Exception:
                n = 0
            if n:
                pivots.append(ws.Name)
                for i in range(n, 0, -1):          # pivot cache trỏ data cũ -> số sai
                    try:
                        ws.PivotTables(i).TableRange2.Clear()
                    except Exception:
                        pass

        olds = template_last_rows or {}
        plan = [(DATA_SHEET, 'AD', layout['data_last_row'], 2)]
        for person, info in layout['per_person'].items():
            plan.append((f'{person}{SUFFIX}', 'N', info['last_row'], 3))
        for sheet, lastcol, newlast, first in plan:
            ws = wb.Worksheets(sheet)
            src = min(10, max(first, newlast))
            try:
                ws.Range(f'A{src}:{lastcol}{src}').Copy()
                ws.Range(f'A{first}:{lastcol}{newlast}').PasteSpecial(Paste=xlPasteFormats)
                app.CutCopyMode = False
            except Exception:
                pass
            old = int(olds.get(sheet, newlast))
            if old > newlast:
                ws.Range(f'A{newlast + 1}:{lastcol}{old + 20}').ClearFormats()

        app.CalculateFullRebuild()
        app.CalculateFull()
        for _ in range(200):
            if app.CalculationState == 0:
                break
            time.sleep(0.1)
        wb.Save()

        results, checks = {}, []
        for person, info in layout['per_person'].items():
            ws = wb.Worksheets(f'{person}{SUFFIX}')
            cc = info['config_col']
            last = info['last_row']
            errs = blanks = 0
            for r in range(3, last + 1):
                for c in (2, 8, 9, 11, 14):
                    v = ws.Cells(r, c).Value
                    if v is None or v == '':
                        blanks += 1
                    if isinstance(v, str) and v.startswith('#'):
                        errs += 1
            results[person] = {
                'time_score': round(float(ws.Cells(9, cc).Value or 0), 2),
                'quality_score': round(float(ws.Cells(9, cc + 1).Value or 0), 2),
                'kpi_score': round(float(ws.Cells(9, cc + 2).Value or 0), 2),
            }
            checks.append({'name': f'{person}: không có ô trống trong vùng dữ liệu',
                           'ok': blanks == 0, 'detail': f'{blanks} ô trống'})
            checks.append({'name': f'{person}: không có lỗi công thức (#N/A, #VALUE!...)',
                           'ok': errs == 0, 'detail': f'{errs} lỗi'})
            nxt = ws.Cells(last + 1, 2).Value
            checks.append({'name': f'{person}: dòng ngay sau vùng dữ liệu phải rỗng',
                           'ok': nxt in (None, ''), 'detail': repr(nxt)})

        master = {}
        if 'Master' in [s.Name for s in wb.Worksheets]:
            ms = wb.Worksheets('Master')
            for r in range(2, 12):
                nm = ms.Cells(r, 2).Value
                if not nm:
                    continue
                try:
                    master[str(nm)] = {
                        'time_score': round(float(ms.Cells(r, 5).Value or 0), 2),
                        'quality_score': round(float(ms.Cells(r, 6).Value or 0), 2),
                        'final_score': round(float(ms.Cells(r, 7).Value or 0), 2),
                        'tier': ms.Cells(r, 8).Value,
                        'money': float(ms.Cells(r, 9).Value or 0),
                    }
                except Exception:
                    pass

        dws = wb.Worksheets(DATA_SHEET)
        sumif = {}
        for i in range(3, 9):
            lab = dws.Cells(i, 31).Value
            if lab:
                sumif[str(lab)] = dws.Cells(i, 32).Value
        for person, info in layout['per_person'].items():
            got = sumif.get(person)
            want = round(sum_t_for(layout, person), 2)
            if got is not None:
                checks.append({'name': f'{person}: SUMIF Data_Tổng khớp tổng giờ phân bổ',
                               'ok': abs(float(got) - want) < 0.01,
                               'detail': f'{got} vs {want}'})
        if pivots:
            checks.append({'name': 'Đã xoá PivotTable cũ còn sót trong template',
                           'ok': True, 'detail': ', '.join(pivots)})

        wb.Close(SaveChanges=True)
        return {'available': True, 'per_person': results, 'master': master,
                'checks': checks, 'sumif': sumif}
    except Exception as exc:
        return {'available': False, 'reason': f'{type(exc).__name__}: {exc}'}
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass


def sum_t_for(layout, person):
    return layout.get('_sum_t', {}).get(person, 0.0)
